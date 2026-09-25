"""Synthetic documents only; no application config, user files, or network."""
from dataclasses import replace
import hashlib
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import time
import unittest
from unittest.mock import patch
import zipfile

from server import material_conversion as conversion
from server.material_conversion import MaterialConversionError, MaterialLimits, convert_material


def pdf_bytes(pages, *, image=False, annotation=False):
    """A complete tiny PDF, built without another parsing/rendering dependency."""
    objects = [b"", b"", b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    page_ids = []
    image_id = None
    if image:
        objects.append(b"<< /Type /XObject /Subtype /Image /Width 1 /Height 1 /ColorSpace /DeviceRGB /BitsPerComponent 8 /Length 3 >>\nstream\n\xff\xff\xff\nendstream")
        image_id = len(objects)
    for text in pages:
        safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = (f"BT /F1 12 Tf 20 100 Td ({safe}) Tj ET\n" if text else "").encode("ascii")
        if image_id:
            stream += b"q 30 0 0 30 20 20 cm /Im0 Do Q\n"
        objects.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"endstream")
        content_id = len(objects)
        annots = b" /Annots [ << /Type /Annot /Subtype /Text /Rect [0 0 1 1] /Contents (Synthetic comment) >> ]" if annotation else b""
        xobject = f" /XObject << /Im0 {image_id} 0 R >>" if image_id else ""
        objects.append((f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] /Resources << /Font << /F1 3 0 R >>{xobject} >> /Contents {content_id} 0 R".encode() + annots + b" >>"))
        page_ids.append(len(objects))
    objects[0] = b"<< /Type /Catalog /Pages 2 0 R >>"
    objects[1] = f"<< /Type /Pages /Count {len(page_ids)} /Kids [{' '.join(f'{i} 0 R' for i in page_ids)}] >>".encode()
    result, offsets = bytearray(b"%PDF-1.4\n"), [0]
    for index, value in enumerate(objects, 1):
        offsets.append(len(result))
        result.extend(f"{index} 0 obj\n".encode() + value + b"\nendobj\n")
    start = len(result)
    result.extend(f"xref\n0 {len(objects)+1}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        result.extend(f"{offset:010d} 00000 n \n".encode())
    result.extend(f"trailer\n<< /Size {len(objects)+1} /Root 1 0 R >>\nstartxref\n{start}\n%%EOF\n".encode())
    return bytes(result)


def pptx_bytes(*, picture=False):
    from pptx import Presentation
    from pptx.util import Inches
    from PIL import Image
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "한국어 슬라이드 원문\n두 번째 문단 <script> & [link](https://invalid.example)"
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(2), Inches(2), Inches(1)).text = "그룹 안 텍스트"
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
    for row, values in enumerate((("항목", "값"), ("표 원문", "A|B\n추가 내용"))):
        for column, value in enumerate(values):
            table.cell(row, column).text = value
    slide.notes_slide.notes_text_frame.text = "발표자 노트의 추가 설명"
    slide2 = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide2.element.set("show", "0")
    slide2.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "숨겨진 슬라이드도 보존"
    if picture:
        raw = io.BytesIO()
        Image.new("RGB", (2, 2), "white").save(raw, format="PNG")
        slide2.shapes.add_picture(io.BytesIO(raw.getvalue()), Inches(2), Inches(2))
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


class MaterialConversionTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.root = Path(self.directory.name)

    def tearDown(self):
        self.directory.cleanup()

    def convert(self, data, kind, **kwargs):
        path = self.root / ("synthetic." + kind)
        path.write_bytes(data)
        before = path.read_bytes()
        result = convert_material(path, kind, **kwargs)
        self.assertEqual(path.read_bytes(), before)
        self.assertEqual(result["sha256"], hashlib.sha256(data).hexdigest())
        self.assertEqual(result["size_bytes"], len(data))
        return result

    def test_pdf_pages_annotations_and_images_retain_provenance_and_original(self):
        result = self.convert(pdf_bytes(["First complete page", "Last complete page"], image=True, annotation=True), "pdf")
        self.assertEqual(result["unit_count"], 2)
        self.assertEqual([unit["source_id"] for unit in result["units"]], ["page:1", "page:2"])
        self.assertIn("First complete page", result["units"][0]["text"])
        self.assertNotIn("Last complete page", result["units"][0]["text"])
        self.assertIn("Last complete page", result["units"][1]["text"])
        self.assertIn("Synthetic comment", result["units"][0]["text"])
        self.assertIn("image_not_extracted", result["warnings"])
        self.assertIn("pdf_table_structure_unverified", result["warnings"])
        self.assertNotIn("ocr_required", result["warnings"])

    def test_scanned_pdf_is_explicitly_pending_ocr_not_fabricated_text(self):
        result = self.convert(pdf_bytes([""], image=True), "pdf")
        self.assertEqual(result["units"][0]["text"], "")
        self.assertIn("ocr_required", result["warnings"])
        self.assertIn("no_extractable_text", result["warnings"])
        self.assertIn("이미지 글자 인식(OCR)", result["markdown"])
        self.assertIn("추출 범위 안내", result["markdown"])

    def test_pptx_all_slides_tables_groups_notes_hidden_content_and_plain_markdown(self):
        result = self.convert(pptx_bytes(picture=True), "pptx")
        self.assertEqual(result["unit_count"], 2)
        first, last = result["units"]
        self.assertEqual([first["source_id"], last["source_id"]], ["slide:1", "slide:2"])
        for text in ("한국어 슬라이드 원문", "두 번째 문단", "그룹 안 텍스트", "표 원문", "A|B", "추가 내용", "발표자 노트의 추가 설명"):
            self.assertIn(text, first["text"])
        self.assertIn("### 발표자 노트", first["markdown"])
        self.assertIn("| 항목 | 값 |", first["markdown"])
        self.assertNotIn("<script>", result["markdown"])
        self.assertNotIn("[link](https://", result["markdown"])
        self.assertIn("숨겨진 슬라이드도 보존", last["text"])
        self.assertIn("hidden_slide", last["warnings"])
        self.assertIn("image_not_extracted", last["warnings"])

    def test_invalid_document_errors_never_include_filename_or_contents(self):
        for kind, data in (("pdf", b"%PDF-secret-private-content"), ("pptx", b"PK\x03\x04secret-private-content"), ("pdf", b"<html>secret</html>")):
            with self.subTest(kind=kind), self.assertRaises(MaterialConversionError) as raised:
                self.convert(data, kind)
            self.assertEqual(str(raised.exception), "invalid_document")
            self.assertNotIn("secret", str(raised.exception))

    def test_limits_reject_whole_document_without_silent_truncation(self):
        for kind, data, limits, code in (
            ("pdf", pdf_bytes(["A", "B"]), MaterialLimits(max_units=1), "unit_limit"),
            ("pptx", pptx_bytes(), MaterialLimits(max_units=1), "unit_limit"),
            ("pdf", pdf_bytes(["This content exceeds the small budget"]), MaterialLimits(max_chars=8), "output_limit"),
            ("pdf", pdf_bytes(["A"]), MaterialLimits(max_file_bytes=8), "size_limit"),
            ("pptx", pptx_bytes(), MaterialLimits(max_archive_entries=2), "archive_limit"),
            ("pptx", pptx_bytes(), MaterialLimits(max_archive_bytes=128), "archive_limit"),
        ):
            with self.subTest(kind=kind, code=code), self.assertRaises(MaterialConversionError) as raised:
                self.convert(data, kind, limits=limits)
            self.assertEqual(raised.exception.code, code)

    def test_zip_traversal_entities_archive_bomb_and_macro_are_rejected(self):
        candidates = [
            ("../outside.xml", b"<a/>", "invalid_document"),
            ("ppt/evil.xml", b'<!DOCTYPE x [<!ENTITY xxe SYSTEM "file:///synthetic-private">]><x>&xxe;</x>', "invalid_document"),
            ("ppt/large.bin", b"X" * 2_000_000, "archive_limit"),
            ("ppt/vbaProject.bin", b"synthetic-macro", "invalid_document"),
        ]
        for name, payload, code in candidates:
            buffer = io.BytesIO(pptx_bytes())
            with zipfile.ZipFile(buffer, "a", compression=zipfile.ZIP_DEFLATED) as archive:
                archive.writestr(name, payload)
            with self.subTest(name=name), self.assertRaises(MaterialConversionError) as raised:
                self.convert(buffer.getvalue(), "pptx")
            self.assertEqual(raised.exception.code, code)
        self.assertFalse((self.root / "outside.xml").exists())

    def test_external_relationship_is_not_followed_and_is_reported(self):
        original = pptx_bytes()
        output = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(original)) as source, zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as dest:
            for entry in source.infolist():
                data = source.read(entry)
                if entry.filename == "ppt/slides/_rels/slide1.xml.rels":
                    data = data.replace(b"</Relationships>", b'<Relationship Id="rIdExternalSynthetic" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink" Target="https://127.0.0.1:1/no-call" TargetMode="External"/></Relationships>')
                dest.writestr(entry.filename, data)
        result = self.convert(output.getvalue(), "pptx")
        self.assertIn("external_links_not_followed", result["warnings"])

    def test_kind_cancel_and_configuration_checks_happen_before_worker(self):
        path = self.root / "missing.pdf"
        for kind, options, code in (("docx", {}, "unsupported_type"), ("pdf", {"cancel": lambda: True}, "cancelled"),
                                    ("pdf", {"limits": MaterialLimits(timeout_seconds=float("nan"))}, "invalid_document")):
            with self.subTest(code=code), patch.object(conversion, "_run_worker", side_effect=AssertionError("No worker allowed")):
                with self.assertRaises(MaterialConversionError) as raised:
                    convert_material(path, kind, **options)
                self.assertEqual(raised.exception.code, code)

    def test_ambient_credentials_proxies_and_pythonpath_not_passed(self):
        with patch.dict(os.environ, {"MINDLOGIC_API_KEY": "synthetic-secret", "GOOGLE_APPLICATION_CREDENTIALS": "synthetic-private",
                                    "HTTPS_PROXY": "synthetic-proxy", "PYTHONPATH": "synthetic-module", "HOME": "synthetic-home"}, clear=True):
            self.assertEqual(conversion._environment(), {})

    def test_timeout_and_cancel_kill_owned_worker_but_not_unrelated_process(self):
        helper = self.root / "synthetic_sleep_worker.py"
        helper.write_text("import sys,time\nsys.stdin.buffer.read()\ntime.sleep(20)\n", encoding="utf-8")
        foreign = subprocess.Popen([sys.executable, "-I", "-c", "import time;time.sleep(30)"], stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        try:
            for cancel, code in ((None, "timeout"), (lambda: True, "cancelled")):
                children = []
                original = subprocess.Popen
                def launch(*args, **kwargs):
                    child = original(*args, **kwargs)
                    children.append(child)
                    return child
                started = time.monotonic()
                with patch.object(conversion, "__file__", str(helper)), patch.object(conversion.subprocess, "Popen", side_effect=launch):
                    with self.assertRaises(MaterialConversionError) as raised:
                        conversion._run_worker(b"synthetic", "pdf", MaterialLimits(timeout_seconds=.5), cancel)
                self.assertEqual(raised.exception.code, code)
                self.assertLess(time.monotonic() - started, 5)
                self.assertTrue(children)
                self.assertTrue(all(child.poll() is not None for child in children))
                self.assertIsNone(foreign.poll())
        finally:
            foreign.terminate()
            foreign.wait(timeout=5)

    @unittest.skipUnless(os.name == "nt", "Windows Job memory enforcement")
    def test_windows_job_enforces_memory_limit_before_parsing(self):
        helper = self.root / "synthetic_memory_worker.py"
        helper.write_text('import sys\nsys.stdin.buffer.read()\ntry:\n x=bytearray(200*1024*1024)\n print("{}")\nexcept MemoryError:\n print(\'{"error":"size_limit"}\')\n', encoding="utf-8")
        with patch.object(conversion, "__file__", str(helper)):
            with self.assertRaises(MaterialConversionError) as raised:
                conversion._run_worker(b"synthetic", "pdf", MaterialLimits(memory_bytes=64*1024*1024), None)
        self.assertEqual(raised.exception.code, "size_limit")

    def test_chart_and_diagram_saved_text_are_retained_with_visual_gaps(self):
        from pptx import Presentation
        from pptx.chart.data import ChartData
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.util import Inches
        presentation = Presentation(io.BytesIO(pptx_bytes()))
        chart = ChartData()
        chart.categories = ["범주 가", "범주 나"]
        chart.add_series("합성 데이터", (3, 5))
        presentation.slides[0].shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(4), Inches(4), Inches(2), chart)
        saved = io.BytesIO();presentation.save(saved)
        output = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(saved.getvalue())) as source, zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as dest:
            for entry in source.infolist():
                data = source.read(entry)
                if entry.filename == "ppt/slides/_rels/slide1.xml.rels":
                    data = data.replace(b"</Relationships>", b'<Relationship Id="rIdSyntheticDiagram" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/diagramData" Target="../diagrams/data1.xml"/></Relationships>')
                dest.writestr(entry.filename, data)
            dest.writestr("ppt/diagrams/data1.xml", '<data xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:t>다이어그램 저장 원문</a:t></data>'.encode())
        result = self.convert(output.getvalue(), "pptx")
        for text in ("범주 가", "범주 나", "합성 데이터", "다이어그램 저장 원문"):
            self.assertIn(text, result["units"][0]["text"])
        self.assertIn("chart_visuals_not_interpreted", result["units"][0]["warnings"])
        self.assertIn("diagram_not_interpreted", result["units"][0]["warnings"])
        self.assertIn("embedded_content_not_extracted", result["warnings"])

    def test_output_pipe_is_bounded_even_for_a_misbehaving_worker(self):
        helper = self.root / "synthetic_noisy_worker.py"
        helper.write_text('import sys\nsys.stdin.buffer.read()\nsys.stdout.buffer.write(b"X"*2000000)\n', encoding="utf-8")
        with patch.object(conversion, "__file__", str(helper)):
            with self.assertRaises(MaterialConversionError) as raised:
                conversion._run_worker(b"synthetic", "pdf", MaterialLimits(max_chars=8, max_units=1), None)
        self.assertEqual(raised.exception.code, "output_limit")

    def test_encrypted_pdf_is_explicitly_rejected(self):
        data = pdf_bytes(["synthetic protected"])
        marker = b"/Root 1 0 R >>"
        encrypted = b"/Root 1 0 R /Encrypt << /Filter /Standard /V 1 /R 2 /Length 40 /O <" + b"00"*32 + b"> /U <" + b"00"*32 + b"> /P -4 >> /ID [<00000000000000000000000000000000><00000000000000000000000000000000>] >>"
        self.assertIn(marker, data)
        with self.assertRaises(MaterialConversionError) as raised:
            self.convert(data.replace(marker, encrypted), "pdf")
        self.assertEqual(raised.exception.code, "encrypted_document")

    def test_worker_result_source_identity_and_provenance_are_revalidated(self):
        data = pdf_bytes(["Only synthetic content"])
        valid = self.convert(data, "pdf")
        for mutate in (lambda doc: doc.update(sha256="0" * 64), lambda doc: doc["units"][0].update(source_id="page:999"),
                       lambda doc: doc["units"][0].update(warnings=["private untrusted error"])):
            result = json.loads(json.dumps(valid))
            mutate(result)
            with self.subTest(mutate=mutate), self.assertRaises(MaterialConversionError) as raised:
                conversion._validate_result(result, data, "pdf", MaterialLimits())
            self.assertEqual(raised.exception.code, "worker_failed")


if __name__ == "__main__":
    unittest.main()
