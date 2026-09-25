"""Bounded, offline PDF/PPTX extraction with page/slide provenance.

Only ``convert_material`` is an application entry point. Parsers run in a fresh
resource-limited child, without application imports, credentials, plugins, URLs,
Office automation, OCR, or shell commands. This is process/resource isolation,
not a general-purpose OS security sandbox. Source paths must come from the
caller's authenticated private upload store, never directly from an HTTP path.
"""
from __future__ import annotations

from dataclasses import asdict, dataclass
import hashlib
import html
import io
import json
import math
import os
from pathlib import Path, PurePosixPath
import re
import signal
import stat
import subprocess
import sys
import threading
import time
from typing import Callable
import zipfile


ERROR_CODES = frozenset({
    "unsupported_type", "invalid_document", "encrypted_document", "size_limit",
    "archive_limit", "unit_limit", "output_limit", "timeout", "cancelled",
    "worker_failed", "dependency_missing", "isolation_unavailable",
})
WARNING_CODES = frozenset({
    "reading_order_unverified", "pdf_table_structure_unverified", "image_not_extracted",
    "vector_graphics_not_interpreted", "ocr_required", "no_extractable_text",
    "font_encoding_unresolved", "control_characters_removed", "diagram_not_interpreted",
    "chart_visuals_not_interpreted", "embedded_content_not_extracted",
    "external_links_not_followed", "hidden_slide", "inherited_content_unverified",
})


WARNING_MESSAGES = {
    "reading_order_unverified": "읽기 순서와 배치를 원본에서 확인해 주세요.",
    "pdf_table_structure_unverified": "PDF 표의 셀 관계는 복원하지 않았습니다. 추출된 글자와 원본 표를 함께 확인해 주세요.",
    "image_not_extracted": "이미지 안의 글자와 시각 자료는 추출하지 않았습니다.",
    "vector_graphics_not_interpreted": "벡터 도형과 도표의 의미는 해석하지 않았습니다.",
    "ocr_required": "추출 가능한 본문이 없습니다. 이미지 글자 인식(OCR) 또는 원본 확인이 필요합니다.",
    "no_extractable_text": "이 페이지 또는 슬라이드에서 본문을 추출하지 못했습니다.",
    "font_encoding_unresolved": "일부 글꼴의 문자 대응을 확인하지 못했습니다.",
    "control_characters_removed": "표시할 수 없는 제어 문자를 제거했습니다.",
    "diagram_not_interpreted": "다이어그램의 시각적 관계는 해석하지 않았습니다.",
    "chart_visuals_not_interpreted": "차트의 저장된 문자와 값만 포함했습니다. 시각적 관계는 원본에서 확인해 주세요.",
    "embedded_content_not_extracted": "첨부 개체 또는 동영상·음성의 내용은 추출하지 않았습니다.",
    "external_links_not_followed": "외부 링크는 열지 않았습니다.",
    "hidden_slide": "숨겨진 슬라이드의 내용도 포함했습니다.",
    "inherited_content_unverified": "슬라이드 마스터·레이아웃에서 상속한 내용은 원본에서 확인해 주세요.",
}


class MaterialConversionError(ValueError):
    def __init__(self, code: str):
        self.code = code if code in ERROR_CODES else "worker_failed"
        super().__init__(self.code)


@dataclass(frozen=True)
class MaterialLimits:
    max_file_bytes: int = 32 * 1024 * 1024
    max_archive_bytes: int = 128 * 1024 * 1024
    max_archive_entries: int = 4096
    max_archive_ratio: int = 100
    max_units: int = 400
    max_chars: int = 1_000_000
    timeout_seconds: float = 60.0
    memory_bytes: int = 768 * 1024 * 1024

    def checked(self):
        maxima = asdict(MaterialLimits())
        for key, value in asdict(self).items():
            if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
                raise MaterialConversionError("invalid_document")
            if key != "timeout_seconds" and not isinstance(value, int):
                raise MaterialConversionError("invalid_document")
            if not 0 < value <= maxima[key]:
                raise MaterialConversionError("invalid_document")
        return self


def _fail(code):
    raise MaterialConversionError(code)


def _text(value, warnings):
    value = str(value).replace("\r\n", "\n").replace("\r", "\n").replace("\v", "\n")
    cleaned = "".join(c for c in value if c in "\n\t" or ord(c) >= 32 and ord(c) != 127)
    if cleaned != value:
        warnings.add("control_characters_removed")
    return cleaned.strip()


def _md(value):
    # Extracted links/images are data, never active Markdown fetches or raw HTML.
    value = html.escape(value, quote=False)
    return re.sub(r"([\\`*_{}\[\]()#!|])", r"\\\1", value)


class _Document:
    def __init__(self, kind, limits):
        self.kind, self.limits, self.units, self.chars = kind, limits, [], 0

    def add(self, blocks, warnings):
        if len(self.units) >= self.limits.max_units:
            _fail("unit_limit")
        text = "\n\n".join(block[0] for block in blocks if block[0])
        markdown = "\n\n".join(block[1] for block in blocks if block[1])
        self.chars += max(len(text), len(markdown))
        if self.chars > self.limits.max_chars:
            _fail("output_limit")
        if not text:
            warnings.add("no_extractable_text")
            if "image_not_extracted" in warnings:
                warnings.add("ocr_required")
        index = len(self.units) + 1
        self.units.append({"index": index, "source_id": f"{'page' if self.kind == 'pdf' else 'slide'}:{index}",
                           "text": text, "markdown": markdown, "warnings": sorted(warnings)})

    def result(self, data):
        if not self.units:
            _fail("invalid_document")
        label = "페이지" if self.kind == "pdf" else "슬라이드"
        sections = []
        for unit in self.units:
            section = f"## {label} {unit['index']}\n\n{unit['markdown']}"
            if unit["warnings"]:
                section += "\n\n### 추출 범위 안내\n\n" + "\n".join(
                    "- " + WARNING_MESSAGES[warning] for warning in unit["warnings"])
            sections.append(section)
        markdown = "\n\n".join(sections)
        if len(markdown) > self.limits.max_chars:
            _fail("output_limit")
        warnings = sorted({warning for unit in self.units for warning in unit["warnings"]})
        return {"kind": self.kind, "sha256": hashlib.sha256(data).hexdigest(), "size_bytes": len(data),
                "unit_count": len(self.units), "units": self.units, "markdown": markdown,
                "warnings": warnings, "extraction_version": "material-v1"}


def _pdf(data, limits):
    from pdfminer.converter import PDFPageAggregator
    from pdfminer.layout import LAParams, LTContainer, LTCurve, LTImage, LTTextContainer
    from pdfminer.pdfdocument import PDFDocument, PDFPasswordIncorrect
    from pdfminer.pdfinterp import PDFPageInterpreter, PDFResourceManager
    from pdfminer.pdfpage import PDFPage
    from pdfminer.pdfparser import PDFParser
    from pdfminer.pdftypes import resolve1
    from pdfminer.utils import decode_text

    result = _Document("pdf", limits)
    stream = io.BytesIO(data)
    try:
        document = PDFDocument(PDFParser(stream))
    except PDFPasswordIncorrect:
        _fail("encrypted_document")
    if document.encryption or not document.is_extractable:
        _fail("encrypted_document")
    resources = PDFResourceManager(caching=False)
    device = PDFPageAggregator(resources, laparams=LAParams(all_texts=True, detect_vertical=True))
    interpreter = PDFPageInterpreter(resources, device)
    try:
        for page in PDFPage.create_pages(document):
            if len(result.units) >= limits.max_units:
                _fail("unit_limit")
            interpreter.process_page(page)
            blocks, warnings = [], {"reading_order_unverified", "pdf_table_structure_unverified"}
            count = 0

            def walk(node, depth=0):
                nonlocal count
                count += 1
                if depth > 64 or count > 200_000:
                    _fail("output_limit")
                if isinstance(node, LTTextContainer):
                    text = _text(node.get_text(), warnings)
                    if re.search(r"\(cid:\d+\)", text):
                        warnings.add("font_encoding_unresolved")
                    if text:
                        blocks.append((text, _md(text)))
                elif isinstance(node, LTImage):
                    warnings.add("image_not_extracted")
                elif isinstance(node, LTCurve):
                    warnings.add("vector_graphics_not_interpreted")
                elif isinstance(node, LTContainer):
                    for child in node:
                        walk(child, depth + 1)
            walk(device.get_result())
            annotations = resolve1(page.annots) if page.annots else []
            if isinstance(annotations, list):
                for annotation in annotations:
                    annotation = resolve1(annotation)
                    if not isinstance(annotation, dict):
                        continue
                    parts = []
                    for key in ("Contents", "T", "V"):
                        value = resolve1(annotation.get(key))
                        if isinstance(value, bytes):
                            parts.append(_text(decode_text(value), warnings))
                        elif isinstance(value, str):
                            parts.append(_text(value, warnings))
                    text = "\n".join(part for part in parts if part)
                    if text:
                        blocks.append((text, "### 주석 / 입력 필드\n\n" + _md(text)))
            result.add(blocks, warnings)
    finally:
        device.close()
        stream.close()
    return result.result(data)


def _zip_preflight(data, limits):
    """Never extract ZIP entries; reject malformed/oversized containers first."""
    from defusedxml import ElementTree as SafeET
    warnings, seen, total = set(), set(), 0
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > limits.max_archive_entries:
            _fail("archive_limit")
        for entry in entries:
            name = entry.filename
            if (not name or "\\" in name or "\x00" in name or ":" in name
                    or name.startswith("/") or ".." in PurePosixPath(name).parts
                    or name.casefold() in seen or stat.S_ISLNK(entry.external_attr >> 16)):
                _fail("invalid_document")
            seen.add(name.casefold())
            if entry.flag_bits & 1:
                _fail("encrypted_document")
            if entry.compress_type not in (zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED):
                _fail("invalid_document")
            total += entry.file_size
            if (total > limits.max_archive_bytes or entry.file_size > limits.max_archive_bytes
                    or entry.file_size > max(1_048_576, entry.compress_size * limits.max_archive_ratio)):
                _fail("archive_limit")
            lower = name.lower()
            if lower.endswith("vbaproject.bin") or lower.endswith("vbadata.xml"):
                _fail("invalid_document")
            if "/embeddings/" in lower or "/media/" in lower and not lower.endswith((".png", ".jpg", ".jpeg", ".gif", ".emf", ".wmf", ".svg", ".tiff", ".tif", ".bmp")):
                warnings.add("embedded_content_not_extracted")
            if lower.endswith((".xml", ".rels")):
                with archive.open(entry) as member:
                    depth, nodes = 0, 0
                    for event, element in SafeET.iterparse(member, events=("start", "end"), forbid_dtd=True, forbid_entities=True, forbid_external=True):
                        if event == "start":
                            depth += 1
                            nodes += 1
                            if depth > 128 or nodes > 200_000:
                                _fail("archive_limit")
                            if element.attrib.get("TargetMode") == "External":
                                warnings.add("external_links_not_followed")
                        else:
                            depth -= 1
                            element.clear()
        if not {"[content_types].xml", "ppt/presentation.xml"}.issubset(seen):
            _fail("invalid_document")
    return warnings


def _pptx(data, limits):
    from defusedxml import ElementTree as SafeET
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    archive_warnings = _zip_preflight(data, limits)
    presentation = Presentation(io.BytesIO(data))
    if len(presentation.slides) > limits.max_units:
        _fail("unit_limit")
    result = _Document("pptx", limits)
    for slide in presentation.slides:
        blocks = []
        warnings = {"reading_order_unverified", "inherited_content_unverified", *archive_warnings}
        if slide.element.get("show") == "0":
            warnings.add("hidden_slide")
        count = 0

        def add_text(value, heading=""):
            text = _text(value, warnings)
            if text:
                blocks.append((text, (heading + "\n\n" if heading else "") + _md(text)))

        def xml_text(blob, *, values=False):
            root = SafeET.fromstring(blob, forbid_dtd=True, forbid_entities=True, forbid_external=True)
            return "\n".join(node.text for node in root.iter()
                             if node.tag.rsplit("}", 1)[-1] in (("t", "v") if values else ("t",)) and node.text)

        def walk(shapes, depth=0):
            nonlocal count
            if depth > 32:
                _fail("output_limit")
            for shape in shapes:
                count += 1
                if count > 20_000:
                    _fail("output_limit")
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    walk(shape.shapes, depth + 1)
                elif shape.has_table:
                    table = shape.table
                    rows = [[_text(cell.text, warnings) for cell in row.cells] for row in table.rows]
                    if rows:
                        plain = "\n".join("\t".join(row) for row in rows)
                        rendered = ["| " + " | ".join(_md(cell).replace("\n", "<br>") for cell in row) + " |" for row in rows]
                        rendered.insert(1, "| " + " | ".join("---" for _ in rows[0]) + " |")
                        blocks.append((plain, "\n".join(rendered)))
                elif shape.has_chart:
                    add_text(xml_text(shape.chart.part.blob, values=True), "### 차트 저장 데이터")
                    warnings.add("chart_visuals_not_interpreted")
                elif shape.has_text_frame:
                    add_text(shape.text_frame.text)
                elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    warnings.add("image_not_extracted")
                    for node in shape.element.iter():
                        if node.tag.endswith("}cNvPr"):
                            add_text("\n".join(node.get(key, "") for key in ("title", "descr")), "### 이미지 대체 설명")
                else:
                    # Connectors, SmartArt and shapes without text remain in the original.
                    warnings.add("diagram_not_interpreted")
        walk(slide.shapes)
        for relation in slide.part.rels.values():
            if relation.is_external:
                warnings.add("external_links_not_followed")
            elif relation.reltype.endswith("/diagramData"):
                add_text(xml_text(relation.target_part.blob), "### 다이어그램 저장 텍스트")
                warnings.add("diagram_not_interpreted")
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            add_text(slide.notes_slide.notes_text_frame.text, "### 발표자 노트")
        result.add(blocks, warnings)
    return result.result(data)


def _extract(data, kind, limits):
    if not data or len(data) > limits.max_file_bytes:
        _fail("size_limit")
    if kind == "pdf" and data.startswith(b"%PDF-"):
        return _pdf(data, limits)
    if kind == "pptx" and data.startswith(b"PK\x03\x04"):
        return _pptx(data, limits)
    _fail("invalid_document")


def _environment():
    # Explicit OS-only values: do not inherit keys, proxies, HOME, PYTHONPATH,
    # cloud settings, user config directories, or model credentials.
    allowed = {"SYSTEMROOT", "WINDIR", "COMSPEC", "LANG", "LC_ALL"}
    return {key: value for key, value in os.environ.items() if key.upper() in allowed}


class _WindowsJob:
    """Own the suspended child tree before any parser or stdin can run."""
    def __init__(self, limits):
        import ctypes
        from ctypes import wintypes as w
        self.ctypes, self.w = ctypes, w
        self.kernel = ctypes.WinDLL("kernel32", use_last_error=True)
        kernel = self.kernel
        for name, args, restype in (
            ("CreateJobObjectW", [ctypes.c_void_p, w.LPCWSTR], w.HANDLE),
            ("SetInformationJobObject", [w.HANDLE, ctypes.c_int, ctypes.c_void_p, w.DWORD], w.BOOL),
            ("AssignProcessToJobObject", [w.HANDLE, w.HANDLE], w.BOOL),
            ("CloseHandle", [w.HANDLE], w.BOOL),
            ("CreateToolhelp32Snapshot", [w.DWORD, w.DWORD], w.HANDLE),
            ("OpenThread", [w.DWORD, w.BOOL, w.DWORD], w.HANDLE),
            ("ResumeThread", [w.HANDLE], w.DWORD),
            ("GetProcessIdOfThread", [w.HANDLE], w.DWORD),
        ):
            function = getattr(kernel, name)
            function.argtypes, function.restype = args, restype
        class Basic(ctypes.Structure):
            _fields_ = [("process_time", ctypes.c_int64), ("job_time", ctypes.c_int64), ("flags", w.DWORD),
                        ("min_working_set", ctypes.c_size_t), ("max_working_set", ctypes.c_size_t),
                        ("active_process_limit", w.DWORD), ("affinity", ctypes.c_size_t),
                        ("priority", w.DWORD), ("scheduling", w.DWORD)]
        class IO(ctypes.Structure):
            _fields_ = [(name, ctypes.c_uint64) for name in ("read_ops", "write_ops", "other_ops", "read_bytes", "write_bytes", "other_bytes")]
        class Extended(ctypes.Structure):
            _fields_ = [("basic", Basic), ("io", IO), ("process_memory", ctypes.c_size_t),
                        ("job_memory", ctypes.c_size_t), ("peak_process_memory", ctypes.c_size_t), ("peak_job_memory", ctypes.c_size_t)]
        self.handle = kernel.CreateJobObjectW(None, None)
        if not self.handle:
            _fail("isolation_unavailable")
        configuration = Extended()
        configuration.basic.flags = 0x2000 | 0x100 | 0x200 | 0x8 | 0x4
        configuration.basic.active_process_limit = 2  # venv redirector and interpreter
        configuration.basic.job_time = int(limits.timeout_seconds * 10_000_000)
        configuration.process_memory = configuration.job_memory = limits.memory_bytes
        if not kernel.SetInformationJobObject(self.handle, 9, ctypes.byref(configuration), ctypes.sizeof(configuration)):
            self.close()
            _fail("isolation_unavailable")

    def attach_resume(self, process):
        ctypes, w, kernel = self.ctypes, self.w, self.kernel
        if not kernel.AssignProcessToJobObject(self.handle, int(process._handle)):
            _fail("isolation_unavailable")
        class Entry(ctypes.Structure):
            _fields_ = [("size", w.DWORD), ("usage", w.DWORD), ("tid", w.DWORD), ("pid", w.DWORD),
                        ("base_priority", w.LONG), ("delta_priority", w.LONG), ("flags", w.DWORD)]
        for name in ("Thread32First", "Thread32Next"):
            function = getattr(kernel, name)
            function.argtypes, function.restype = [w.HANDLE, ctypes.POINTER(Entry)], w.BOOL
        snapshot = kernel.CreateToolhelp32Snapshot(4, 0)
        if snapshot == ctypes.c_void_p(-1).value:
            _fail("isolation_unavailable")
        try:
            entry, threads = Entry(), []
            entry.size = ctypes.sizeof(entry)
            available = kernel.Thread32First(snapshot, ctypes.byref(entry))
            while available:
                if entry.pid == process.pid:
                    threads.append(entry.tid)
                entry.size = ctypes.sizeof(entry)
                available = kernel.Thread32Next(snapshot, ctypes.byref(entry))
            if len(threads) != 1:
                _fail("isolation_unavailable")
            thread = kernel.OpenThread(0x0802, False, threads[0])
            if not thread:
                _fail("isolation_unavailable")
            try:
                if kernel.GetProcessIdOfThread(thread) != process.pid or kernel.ResumeThread(thread) != 1:
                    _fail("isolation_unavailable")
            finally:
                kernel.CloseHandle(thread)
        finally:
            kernel.CloseHandle(snapshot)

    def close(self):
        if self.handle:
            self.kernel.CloseHandle(self.handle)
            self.handle = None


def _run_worker(data, kind, limits, cancel):
    process, job, output = None, None, bytearray()
    overflow, reader_done = threading.Event(), threading.Event()
    # text + per-unit Markdown + whole Markdown + JSON escaping, bounded above.
    max_response = limits.max_chars * 20 + limits.max_units * 4096 + 65536
    packet = json.dumps({"kind": kind, "limits": asdict(limits)}, separators=(",", ":")).encode() + b"\n" + data
    command = [sys.executable, "-I", "-B", str(Path(__file__).resolve()), "--worker"]
    deadline = time.monotonic() + limits.timeout_seconds
    try:
        if os.name == "nt":
            job = _WindowsJob(limits)
        process = subprocess.Popen(command, stdin=subprocess.PIPE, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL,
            close_fds=True, env=_environment(), cwd=Path(sys.executable).parent,
            creationflags=(0x4 | subprocess.CREATE_NO_WINDOW) if os.name == "nt" else 0,
            start_new_session=os.name != "nt")
        if job:
            job.attach_resume(process)

        def write():
            try:
                process.stdin.write(packet)
            except (OSError, ValueError):
                pass
            finally:
                try:
                    process.stdin.close()
                except (OSError, ValueError):
                    pass

        def read():
            try:
                while block := process.stdout.read(65536):
                    if len(output) + len(block) > max_response:
                        overflow.set()
                        break
                    output.extend(block)
            finally:
                reader_done.set()
        writer = threading.Thread(target=write, daemon=True)
        reader = threading.Thread(target=read, daemon=True)
        writer.start()
        reader.start()
        while process.poll() is None or not reader_done.is_set():
            if cancel is not None and cancel():
                _fail("cancelled")
            if overflow.is_set():
                _fail("output_limit")
            if time.monotonic() >= deadline:
                _fail("timeout")
            reader_done.wait(0.02) if process.poll() is not None else time.sleep(0.02)
        writer.join(timeout=1)
        reader.join(timeout=1)
        if overflow.is_set():
            _fail("output_limit")
        if process.returncode != 0:
            _fail("worker_failed")
        try:
            response = json.loads(output)
        except (ValueError, UnicodeError):
            _fail("worker_failed")
        if not isinstance(response, dict):
            _fail("worker_failed")
        if "error" in response:
            _fail(response["error"] if isinstance(response["error"], str) else "worker_failed")
        return response
    except MaterialConversionError:
        raise
    except (OSError, ValueError):
        raise MaterialConversionError("worker_failed") from None
    finally:
        if job:
            job.close()  # only this conversion's owned process tree
        if process is not None:
            if os.name != "nt" and process.poll() is None:
                try:
                    os.killpg(process.pid, signal.SIGKILL)
                except ProcessLookupError:
                    pass
            elif process.poll() is None:
                process.kill()  # includes suspended process if job assignment failed
            process.wait(timeout=5)
            for pipe in (process.stdin, process.stdout):
                if pipe is not None:
                    pipe.close()


def _validate_result(result, data, kind, limits):
    if (set(result) != {"kind", "sha256", "size_bytes", "unit_count", "units", "markdown", "warnings", "extraction_version"}
            or result["kind"] != kind or result["sha256"] != hashlib.sha256(data).hexdigest()
            or result["size_bytes"] != len(data) or result["extraction_version"] != "material-v1"
            or not isinstance(result["units"], list) or not 0 < len(result["units"]) <= limits.max_units
            or result["unit_count"] != len(result["units"])):
        _fail("worker_failed")
    doc = _Document(kind, limits)
    for index, unit in enumerate(result["units"], 1):
        if (not isinstance(unit, dict) or set(unit) != {"index", "source_id", "text", "markdown", "warnings"}
                or unit["index"] != index or unit["source_id"] != f"{'page' if kind == 'pdf' else 'slide'}:{index}"
                or not isinstance(unit["text"], str) or not isinstance(unit["markdown"], str)
                or not isinstance(unit["warnings"], list) or any(not isinstance(w, str) or w not in WARNING_CODES for w in unit["warnings"])):
            _fail("worker_failed")
        doc.add([(unit["text"], unit["markdown"])], set(unit["warnings"]))
    expected = doc.result(data)
    if expected != result:
        _fail("worker_failed")
    return result


def convert_material(source: Path, kind: str, *, limits: MaterialLimits | None = None,
                     cancel: Callable[[], bool] | None = None) -> dict:
    """Extract a private source file. Caller keeps originals and enforces ownership.

    Warnings are coverage gaps, not a claim of complete visual understanding.
    OCR/image-only units are returned empty with ``ocr_required``; callers must
    keep them pending for review rather than label them a complete conversion.
    """
    if kind not in ("pdf", "pptx"):
        _fail("unsupported_type")
    limits = (limits or MaterialLimits()).checked()
    if cancel is not None and cancel():
        _fail("cancelled")
    try:
        source = Path(source)
        if source.is_symlink() or not stat.S_ISREG(source.stat().st_mode):
            _fail("invalid_document")
        with source.open("rb") as stream:
            data = stream.read(limits.max_file_bytes + 1)
    except (OSError, ValueError):
        raise MaterialConversionError("invalid_document") from None
    if not data or len(data) > limits.max_file_bytes:
        _fail("size_limit")
    result = _run_worker(data, kind, limits, cancel)
    return _validate_result(result, data, kind, limits)


def _worker_main():
    import logging
    logging.disable(logging.CRITICAL)
    try:
        config = json.loads(sys.stdin.buffer.readline(4096))
        kind, limits = config["kind"], MaterialLimits(**config["limits"]).checked()
        if kind not in ("pdf", "pptx"):
            _fail("unsupported_type")
        if os.name != "nt":
            import resource
            resource.setrlimit(resource.RLIMIT_AS, (limits.memory_bytes, limits.memory_bytes))
            seconds = max(1, math.ceil(limits.timeout_seconds))
            resource.setrlimit(resource.RLIMIT_CPU, (seconds, seconds))
            resource.setrlimit(resource.RLIMIT_NOFILE, (64, 64))
            resource.setrlimit(resource.RLIMIT_FSIZE, (0, 0))

        def no_external_io(event, args):
            if event.startswith("socket.") or event in ("subprocess.Popen", "os.system", "os.exec", "os.posix_spawn"):
                _fail("invalid_document")
        sys.addaudithook(no_external_io)
        data = sys.stdin.buffer.read(limits.max_file_bytes + 1)
        result = _extract(data, kind, limits)
    except MaterialConversionError as error:
        result = {"error": error.code}
    except ImportError:
        result = {"error": "dependency_missing"}
    except MemoryError:
        result = {"error": "size_limit"}
    except BaseException:
        result = {"error": "invalid_document"}
    sys.stdout.buffer.write(json.dumps(result, ensure_ascii=False, separators=(",", ":")).encode("utf-8"))
    sys.stdout.buffer.flush()


if __name__ == "__main__":
    if sys.argv[1:] != ["--worker"]:
        raise SystemExit(2)
    _worker_main()
